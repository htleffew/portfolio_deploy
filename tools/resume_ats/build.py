"""Build the styled, ATS-parseable resume (HTML->PDF via Chromium, and DOCX) from one content model."""
import re, html, os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))

# ---------- content model (parsed from the original resume text) ----------
lines = open(os.path.join(HERE, 'resume.txt')).read().replace('\f', '').split('\n')
sections = {}; cur = None
for l in lines[6:]:
    if re.fullmatch(r'[A-Z][A-Z ]+', l.strip()) and not l.startswith(' '):
        cur = l.strip(); sections[cur] = []; continue
    if cur: sections[cur].append(l)

def join(parts):
    out = ''
    for p in parts:
        p = p.strip()
        out = p if not out else (out + p if out.endswith('-') else out + ' ' + p)
    return out

def split_entries(ls, startfn):
    ents, c = [], []
    for l in ls:
        if not l.strip(): continue
        if c and startfn(c[-1], l): ents.append(join(c)); c = []
        c.append(l)
    if c: ents.append(join(c))
    return ents

newent = lambda prev, l: prev.rstrip().endswith('.')

NAME = 'Heather Leffew, PhD'
TITLE = ['Behavioral Scientist', 'Psychological Text Analysis', 'Experimentation and Research Leadership']
TAGLINE = 'Applying the science of human behavior to language, decisions, and measurable business results.'
CONTACT = [('htmleffew@gmail.com', 'mailto:htmleffew@gmail.com'),
           ('DrHeatherLeffew.com', 'https://drheatherleffew.com/'),
           ('linkedin.com/in/heathertleffew', 'https://www.linkedin.com/in/heathertleffew')]
# research links, recovered from the original PDF's link annotations
SITE = 'https://drheatherleffew.com/'
RESEARCH_LINKS = {'Instrumental and Affective': 'doctoral-dissertation/', 'LIWC and the TAT': 'diagnostic-thematic-apperception-liwc/',
                  'Implicit Power Drives': 'implicit-power-drives/', 'Complaint Text': 'complaint-text-monitor/',
                  "The Helper's Language": 'helper-language-signal/', 'UX Principles': 'ux-livestreamed-violence/',
                  'The AI Productivity': 'ai-productivity-j-curve/', 'Pathologizing Without Warrant': 'claude-lcr-analysis/',
                  'The Somatic Deficit': 'somatic-deficit/'}
def research_link(title):
    return next(SITE + v for k, v in RESEARCH_LINKS.items() if title.startswith(k))
SUMMARY = join([l for l in sections['SUMMARY'] if l.strip()])

jobre = re.compile(r'^(\S.*?)\s{3,}((?:\w{3} \d{4}) to (?:Present|\w{3} \d{4}))$')
JOBS = []
for l in sections['EMPLOYMENT HISTORY']:
    m = jobre.match(l)
    if m: JOBS.append({'title': m.group(1), 'date': m.group(2), 'ls': []})
    elif JOBS: JOBS[-1]['ls'].append(l)
for j in JOBS:
    body = [l for l in j.pop('ls') if l.strip()]
    j['company'] = body.pop(0).strip().replace(' | ', ', ') if (not body[0].startswith('   ') and len(body[0]) < 80) else None
    k = next((n for n, l in enumerate(body) if l.startswith('   ')), len(body))
    j['intro'] = join(body[:k]) if k else None
    j['bullets'] = split_entries(body[k:], newent)

ed = [l.strip() for l in sections['EDUCATION'] if l.strip()]
EDUCATION = [(ed[k], ed[k + 1].replace(' | ', ', ')) for k in range(0, len(ed), 2)]
RESEARCH = []
for e in split_entries(sections['SELECTED RESEARCH'], newent):
    m = re.match(r'^(.+?\.)\s(.*)$', e.replace('(dissertation).', '(dissertation)#').replace('(preprint).', '(preprint)#').replace('#', '.'))
    RESEARCH.append((m.group(1), m.group(2), research_link(m.group(1))))

slabels = 'AI Evaluation & Safety|Measurement & Research Methods|Machine Learning & Modeling|MLOps, Platform & Infrastructure|Governance, Safety & Compliance|Technical Stack'
sl = [l for l in open(os.path.join(HERE, 'skills.txt')).read().split('\n') if l.strip()]
SKILLS = [list(re.match(r'^([^:]+): (.*)$', e).groups()) for e in split_entries(sl, lambda p, l: re.match(r'^(%s): ' % slabels, l))]
for x in SKILLS:
    if x[0] == 'Measurement & Research Methods':
        x[1] += (", Minimum Detectable Effect (MDE), Difference-in-Differences, Within-Subjects Designs, Preregistration, "
                 "Offline Replay Experiments, FDR-Controlled Subgroup Analysis, Break-Even and ROI Modeling, Factor Analysis, "
                 "Cohen's Kappa, PABAK, Adverse-Impact Analysis, LIWC, Hand-Coded Content Analysis, Text Classification, "
                 "LDA Topic Modeling, Collocation Analysis (logDice, nMPI), Discourse Analysis, Co-occurrence Network Analysis")
    if x[0] == 'Technical Stack':
        x[1] = x[1].replace('SQL,', 'SQL, R, SAS, SPSS, Excel,', 1)
SKILLS.insert([x[0] for x in SKILLS].index('Measurement & Research Methods') + 1,
              ['Behavioral Science & Qualitative Methods',
               "Pathway to Violence, Warning-Behavior Typologies, Habit Formation, Choice Friction in Interaction Design, Hick's Law, Jakob's Law, "
               "von Restorff Effect, Ironic Process Theory, Kohlberg Moral Development, Semi-Structured Interviews, User Research, "
               "Behavioral Observation, Clinical Assessment, Forensic Assessment, Structured Threat Assessment"])

# ---------- HTML (for PDF) ----------
NAVY, INK, GREY = '#25365e', '#1f2328', '#5f6670'

def e(t):
    # keep hyphenated compounds on one line: extractors drop a hyphen at a line break
    t = html.escape(t)
    return re.sub(r'(\w+(?:-\w+)+)', r'<span class="nw">\1</span>', t)

SEP = '<span class="sep"> | </span>'
F = os.path.join(HERE, 'fonts')
css = f"""
@font-face{{font-family:Lora;src:url(file://{F}/lora-latin-400-normal.woff2);font-weight:400}}
@font-face{{font-family:Lora;src:url(file://{F}/lora-latin-400-italic.woff2);font-weight:400;font-style:italic}}
@font-face{{font-family:Lora;src:url(file://{F}/lora-latin-600-normal.woff2);font-weight:600}}
@font-face{{font-family:Lora;src:url(file://{F}/lora-latin-700-normal.woff2);font-weight:700}}
@font-face{{font-family:Lora;src:url(file://{F}/lora-latin-600-italic.woff2);font-weight:600;font-style:italic}}
@font-face{{font-family:Playfair;src:url(file://{F}/playfair-display-latin-600-normal.woff2);font-weight:600}}
@page{{size:Letter;margin:0.5in 0.65in}}
body{{font:9.6pt/1.42 Lora,Georgia,serif;color:{INK};margin:0}}
p{{margin:0 0 3.5pt;break-inside:avoid;orphans:3;widows:3}}
.nw{{white-space:nowrap}}
header{{text-align:center;margin-bottom:6pt}}
.name{{font:600 23pt/1.1 Playfair,Georgia,serif;color:{NAVY};margin:0 0 3pt}}
.title{{font-size:10pt;margin:0 0 1pt}}
.tag{{font-style:italic;color:{GREY};margin:0 0 3pt}}
.contact{{color:{NAVY}}}
.sep{{color:#9aa1ab}}
h2{{font:600 10pt/1 Lora,serif;color:{NAVY};text-transform:uppercase;letter-spacing:.06em;margin:15pt 0 7pt;padding-bottom:3pt;border-bottom:1.2px solid {NAVY};break-after:avoid}}
.job{{font-weight:700;font-size:10.4pt;margin:11pt 0 0;break-after:avoid}}
h2+.job{{margin-top:0}}
.sub{{margin:0 0 4pt;break-after:avoid}}
a{{color:inherit;text-decoration:none}}
.contact a{{border-bottom:1px solid #c5cbd6}}
.job b{{font-weight:700}}
.co{{color:{NAVY};font-weight:600}}
.dt{{color:{GREY};font-style:italic}}
.b{{padding-left:12pt;text-indent:-12pt;margin-bottom:3pt}}
.bul{{color:{NAVY};display:inline-block;width:12pt;text-indent:0}}
.lab{{font-weight:700;color:{NAVY}}}
.rt{{font-weight:600;font-style:italic;color:{NAVY}}}
.ed{{margin-bottom:5pt}}
"""
H = [f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{NAME} - Resume</title><style>{css}</style></head><body>',
     '<header>',
     f'<p class="name">{NAME}</p>',
     f'<p class="title">{SEP.join(e(t) for t in TITLE)}</p>',
     f'<p class="tag">{e(TAGLINE)}</p>',
     '<p class="contact">' + SEP.join(f'<a href="{u}">{html.escape(c)}</a>' for c, u in CONTACT) + '</p>',
     '</header>']
bullet = lambda t: f'<p class="b"><span class="bul">&bull;</span>{e(t)}</p>'
H += ['<h2>Summary</h2>', f'<p>{e(SUMMARY)}</p>', '<h2>Employment History</h2>']
for j in JOBS:
    H.append(f'<p class="job">{e(j["title"])}</p>')
    sub = ([f'<span class="co">{e(j["company"])}</span>'] if j['company'] else []) + [f'<span class="dt nw">{e(j["date"])}</span>']
    H.append(f'<p class="sub">{SEP.join(sub)}</p>')
    if j['intro']: H.append(f'<p>{e(j["intro"])}</p>')
    H += [bullet(b) for b in j['bullets']]
H.append('<h2>Education</h2>')
for deg, sch in EDUCATION:
    H.append(f'<p class="ed"><b>{e(deg)}</b><br><span class="dt">{e(sch)}</span></p>')
H.append('<h2>Selected Research</h2>')
for t, d, u in RESEARCH:
    H.append(f'<p class="b"><span class="bul">&bull;</span><a class="rt" href="{u}">{e(t)}</a> {e(d)}</p>')
H.append('<h2>Skills</h2>')
for lab, txt in SKILLS:
    H.append(f'<p><span class="lab">{e(lab)}:</span> {e(txt)}</p>')
H.append('</body></html>')
open(os.path.join(HERE, 'resume_ats.html'), 'w').write('\n'.join(H))

# ---------- DOCX ----------
def rgb(h): return RGBColor.from_string(h.lstrip('#').upper())
SERIF, DISPLAY = 'Georgia', 'Georgia'  # universally installed; Lora/Playfair are not
doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Inches(0.65); sec.top_margin = sec.bottom_margin = Inches(0.5)
st = doc.styles['Normal']; st.font.name = SERIF; st.font.size = Pt(9); st.font.color.rgb = rgb(INK)
st.element.rPr.rFonts.set(qn('w:eastAsia'), SERIF)
st.paragraph_format.line_spacing = 1.2

def P(after=3.5, before=0, align=None, keep_next=False, hang=None):
    p = doc.add_paragraph(); f = p.paragraph_format
    f.space_after = Pt(after); f.space_before = Pt(before); f.keep_together = True; f.keep_with_next = keep_next
    if align: p.alignment = align
    if hang: f.left_indent = Inches(hang); f.first_line_indent = Inches(-hang); f.tab_stops.add_tab_stop(Inches(hang))
    return p

def R(p, t, bold=False, italic=False, color=None, size=None, font=None):
    r = p.add_run(t); r.bold = bold; r.italic = italic
    if color: r.font.color.rgb = rgb(color)
    if size: r.font.size = Pt(size)
    if font: r.font.name = font
    return r

def link(p, text, url, **kw):
    rid = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    h = OxmlElement('w:hyperlink'); h.set(qn('r:id'), rid)
    r = R(p, text, **kw); h.append(r._r); p._p.append(h)

def sep_join(p, items):
    for i, (t, kw) in enumerate(items):
        if i: R(p, ' | ', color='#9aa1ab')
        R(p, t, **kw)

def heading(t):
    p = P(after=6, before=13, keep_next=True)
    r = R(p, t.upper(), bold=True, color=NAVY, size=10)
    rPr = r._r.get_or_add_rPr(); sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '12'); rPr.append(sp)
    bdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '8'), ('space', '2'), ('color', NAVY.lstrip('#'))): b.set(qn('w:' + k), v)
    bdr.append(b); p._p.get_or_add_pPr().append(bdr)

C = WD_ALIGN_PARAGRAPH.CENTER
R(P(after=2, align=C), NAME, color=NAVY, size=22, font=DISPLAY)
sep_join(P(after=1, align=C), [(t, {'size': 9.5}) for t in TITLE])
R(P(after=2, align=C), TAGLINE, italic=True, color=GREY)
p = P(after=4, align=C)
for i, (c, u) in enumerate(CONTACT):
    if i: R(p, ' | ', color='#9aa1ab')
    link(p, c, u, color=NAVY)
heading('Summary'); R(P(), SUMMARY)
heading('Employment History')
for j in JOBS:
    R(P(before=9, after=0, keep_next=True), j['title'], bold=True, size=10)
    items = ([(j['company'], {'bold': True, 'color': NAVY})] if j['company'] else []) + [(j['date'], {'italic': True, 'color': GREY})]
    sep_join(P(after=3, keep_next=True), items)
    if j['intro']: R(P(), j['intro'])
    for b in j['bullets']:
        p = P(hang=0.15); R(p, '•\t', color=NAVY); R(p, b)
heading('Education')
for deg, sch in EDUCATION:
    R(P(after=0, keep_next=True), deg, bold=True); R(P(after=5), sch, italic=True, color=GREY)
heading('Selected Research')
for t, d, u in RESEARCH:
    p = P(hang=0.15); R(p, '•\t', color=NAVY); link(p, t, u, bold=True, italic=True, color=NAVY); R(p, ' ' + d)
heading('Skills')
for lab, txt in SKILLS:
    p = P(); R(p, lab + ': ', bold=True, color=NAVY); R(p, txt)
doc.core_properties.title = NAME + ' - Resume'; doc.core_properties.author = 'Heather Leffew'
doc.save(os.path.join(HERE, 'Heather_Leffew_Resume_ATS.docx'))
